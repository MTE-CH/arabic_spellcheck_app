"""
Core logic for the Arabic (MSA) wrong-word detector.

Two steps are applied to every word:
  1. CONTEXT   - CAMeLBERT-MSA (masked language model) scores how likely the word is in
                 its sentence. Words below a very low threshold are flagged.
  2. DICTIONARY - words not flagged in step 1 are looked up in the AyaSpell (Hunspell)
                 Arabic dictionary. Words it rejects are flagged too.

Nothing is rewritten and no corrections are suggested.
"""
from __future__ import annotations

import io
import json
import os
import random
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

import numpy as np

MODEL_NAME = "CAMeL-Lab/bert-base-arabic-camelbert-msa"
PROJECT_DIR = Path(__file__).resolve().parent
PROFILES_DIR = PROJECT_DIR / "profiles"
CUSTOM_WORDS_PATH = PROJECT_DIR / "custom_words.txt"

# Used only when a profile has no calibrated value. It is a starting guess, not a measurement.
FALLBACK_THRESHOLD = -12.0

# --------------------------------------------------------------------------------------
# Text handling
# --------------------------------------------------------------------------------------
ARABIC_LETTER = re.compile(r"[ء-ي]")
LATIN_OR_DIGIT = re.compile(r"[A-Za-z0-9]")
# diacritics, tatweel, superscript alef, zero-width and bidi control characters
STRIP_CHARS = re.compile(
    r"[ً-ْٰـ​-‏‪-‮⁦-⁩﻿]"
)
# a word may contain diacritics in the ORIGINAL text, so they are part of the token there
TOKEN_RE = re.compile(
    r"[\wؐ-ًؚ-ٰٟۖ-ۭـ]+|[^\w\s]"
)
SENTENCE_RE = re.compile(r"[^.!?؟؛\n]+[.!?؟؛]*")


def normalize(text: str) -> str:
    """NFKC (repairs Arabic presentation forms from PDFs) + remove diacritics/tatweel.

    Hamza, alef and taa marbuta are deliberately kept: mistakes there are what we detect.
    """
    return STRIP_CHARS.sub("", unicodedata.normalize("NFKC", text))


def is_arabic_word(w: str) -> bool:
    return bool(ARABIC_LETTER.search(w)) and not LATIN_OR_DIGIT.search(w)


@dataclass
class Token:
    raw: str
    norm: str
    start: int
    end: int


def tokenize(sentence: str, offset: int = 0) -> List[Token]:
    toks = []
    for m in TOKEN_RE.finditer(sentence):
        raw = m.group()
        norm = normalize(raw)
        if norm:
            toks.append(Token(raw, norm, offset + m.start(), offset + m.end()))
    return toks


def iter_sentences(text: str):
    for m in SENTENCE_RE.finditer(text):
        if m.group().strip():
            yield m.start(), m.group()


# --------------------------------------------------------------------------------------
# Reading uploaded files
# --------------------------------------------------------------------------------------
def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1256"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith((".txt", ".md")):
        return decode_bytes(data)
    if name.endswith(".docx"):
        import docx  # python-docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    raise ValueError(f"Unsupported file type: {filename} (use .txt, .docx or .pdf)")


# --------------------------------------------------------------------------------------
# Dictionary (AyaSpell via Hunspell)
# --------------------------------------------------------------------------------------
def find_dictionary_path() -> Optional[str]:
    candidates = [
        os.environ.get("ARABIC_DICT_PATH"),
        str(PROJECT_DIR / "dict" / "ar"),
        "/usr/share/hunspell/ar",
        "/usr/share/myspell/ar",
        "/usr/share/myspell/dicts/ar",
    ]
    for base in candidates:
        if base and os.path.exists(base + ".dic") and os.path.exists(base + ".aff"):
            return base
    return None


def load_dictionary(base: Optional[str] = None):
    from spylls.hunspell import Dictionary

    base = base or find_dictionary_path()
    if not base:
        raise FileNotFoundError(
            "Arabic Hunspell dictionary not found. Put ar.dic and ar.aff in the 'dict' "
            "folder, or set ARABIC_DICT_PATH to the path without extension."
        )
    return Dictionary.from_files(base)


def prefix_variants(w: str) -> set:
    """Forms of w with common attached prefixes removed (و/ف + ب/ل/ك + ال).

    Used only for the custom-words list, because AyaSpell already handles affixes itself.
    """
    out = {w}
    for conj in ("", "و", "ف"):
        if conj and not w.startswith(conj):
            continue
        r1 = w[len(conj):]
        for prep in ("", "ب", "ل", "ك"):
            if prep and not r1.startswith(prep):
                continue
            r2 = r1[len(prep):]
            out.add(r2)
            if r2.startswith("ال"):
                out.add(r2[2:])
            if prep == "ل" and r2.startswith("ل"):  # للمدير = ل + المدير
                out.add(r2[1:])
    return {v for v in out if len(v) >= 2}


class WordChecker:
    """Dictionary check with a custom-words list (names, places, domain terms)."""

    def __init__(self, dic, custom=(), cache: Optional[dict] = None):
        self.dic = dic
        self.custom = {normalize(c).strip() for c in custom if c.strip()}
        self._cache = cache if cache is not None else {}

    def is_custom(self, w: str) -> bool:
        if not self.custom:
            return False
        return any(v in self.custom for v in prefix_variants(w))

    def _lookup(self, w: str) -> bool:
        r = self._cache.get(w)
        if r is None:
            try:
                r = bool(self.dic.lookup(w))
            except Exception:
                r = True  # if we cannot check a word, do not flag it
            self._cache[w] = r
        return r

    def ok(self, w: str) -> bool:
        return self.is_custom(w) or self._lookup(w)


def load_custom_words(path: Path = CUSTOM_WORDS_PATH) -> List[str]:
    if not Path(path).exists():
        return []
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    return [l.strip() for l in lines if l.strip() and not l.strip().startswith("#")]


# --------------------------------------------------------------------------------------
# Context model (CAMeLBERT-MSA as a masked language model)
# --------------------------------------------------------------------------------------
class ContextScorer:
    """Scores each word by the mean log-probability of its pieces when ALL of them are masked."""

    def __init__(self, tokenizer, model, device: str = "cpu", max_length: int = 256):
        self.tok = tokenizer
        self.model = model
        self.device = device
        self.max_length = max_length

    @classmethod
    def from_pretrained(cls, name: str = MODEL_NAME, device: Optional[str] = None):
        import torch
        from transformers import AutoModelForMaskedLM, AutoTokenizer

        device = device or ("cuda" if torch.cuda.is_available() else "cpu")
        tok = AutoTokenizer.from_pretrained(name)
        model = AutoModelForMaskedLM.from_pretrained(name).to(device).eval()
        return cls(tok, model, device)

    def score_words(self, words: List[str], batch_size: int = 32) -> Dict[int, float]:
        """Return {index: mean log-prob} for each Arabic word in `words`."""
        import torch

        if not words:
            return {}
        enc = self.tok(
            words,
            is_split_into_words=True,
            return_tensors="pt",
            truncation=True,
            max_length=self.max_length,
        )
        ids = enc["input_ids"][0].to(self.device)
        word_ids = enc.word_ids(0)
        spans: Dict[int, List[int]] = {}
        for pos, w in enumerate(word_ids):
            if w is not None:
                spans.setdefault(w, []).append(pos)

        targets = [w for w in spans if is_arabic_word(words[w])]
        out: Dict[int, float] = {}
        with torch.no_grad():
            for i in range(0, len(targets), batch_size):
                chunk = targets[i:i + batch_size]
                inp = ids.unsqueeze(0).repeat(len(chunk), 1)
                for r, w in enumerate(chunk):
                    inp[r, spans[w]] = self.tok.mask_token_id
                logp = self.model(input_ids=inp).logits.log_softmax(-1)
                for r, w in enumerate(chunk):
                    p = torch.tensor(spans[w], device=self.device)
                    out[w] = logp[r][p, ids[p]].mean().item()
        return out


def load_scorer() -> ContextScorer:
    return ContextScorer.from_pretrained()


# --------------------------------------------------------------------------------------
# Checking text
# --------------------------------------------------------------------------------------
@dataclass
class Flag:
    start: int
    end: int
    word: str
    reason: str            # "context" or "dictionary"
    score: Optional[float]


def analyze(
    text: str,
    scorer,
    checker: WordChecker,
    threshold: float,
    chunk_words: int = 60,
    progress: Optional[Callable[[float], None]] = None,
) -> Tuple[List[Flag], int]:
    """Apply the two steps. Returns (flags, number_of_arabic_words_checked).

    Offsets in each Flag refer to the ORIGINAL text, so the caller can highlight it.
    Words in the custom list are never flagged.
    """
    sentences = list(iter_sentences(text))
    flags: List[Flag] = []
    n_words = 0
    for si, (off, sent) in enumerate(sentences):
        toks = tokenize(sent, off)
        for i in range(0, len(toks), chunk_words):
            chunk = toks[i:i + chunk_words]
            if not any(is_arabic_word(t.norm) for t in chunk):
                continue
            scores = scorer.score_words([t.norm for t in chunk])
            for j, t in enumerate(chunk):
                if not is_arabic_word(t.norm):
                    continue
                n_words += 1
                if checker.is_custom(t.norm):
                    continue
                sc = scores.get(j)
                if sc is not None and sc < threshold:
                    reason = "context"                       # step 1
                elif not checker._lookup(t.norm):
                    reason = "dictionary"                    # step 2 (only for words not flagged in step 1)
                else:
                    continue
                flags.append(Flag(t.start, t.end, t.raw, reason, sc))
        if progress:
            progress((si + 1) / max(len(sentences), 1))
    return flags, n_words


# --------------------------------------------------------------------------------------
# Calibration
# --------------------------------------------------------------------------------------
@dataclass
class CalibrationResult:
    threshold: float
    percentile: float
    scores: np.ndarray
    n_sentences: int
    n_words: int
    stats: dict = field(default_factory=dict)

    @property
    def note(self) -> str:
        return reliability_note(self.n_words, self.percentile)


def reliability_note(n_words: int, percentile: float) -> str:
    tail = n_words * percentile / 100.0
    if tail < 30:
        return (
            f"Only about {tail:.0f} words fall below the cutoff, so the threshold is unreliable. "
            "Use more text or a higher percentile."
        )
    if tail < 100:
        return f"About {tail:.0f} words define the cutoff. Usable, but more text would be steadier."
    return f"About {tail:.0f} words define the cutoff. This is a stable estimate."


def prepare_calibration_sentences(
    text: str,
    checker: WordChecker,
    max_sentences: int = 2000,
    min_words: int = 6,
    max_words: int = 40,
    max_reject_ratio: float = 0.25,
    seed: int = 0,
    progress: Optional[Callable[[float], None]] = None,
) -> Tuple[List[str], dict]:
    """Pick a random sample of usable sentences from a book.

    Sentences are dropped if they are too short/long, mostly non-Arabic, or look garbled
    (a large share of their words is rejected by the dictionary, typical of bad OCR/PDF text).
    """
    candidates: List[str] = []
    for _, sent in iter_sentences(text):
        s = normalize(sent).strip()
        words = [m.group() for m in TOKEN_RE.finditer(s) if re.match(r"\w", m.group())]
        ar = [w for w in words if is_arabic_word(w)]
        if not (min_words <= len(ar) <= max_words):
            continue
        if len(ar) / max(len(words), 1) < 0.9:
            continue
        candidates.append(s)

    random.Random(seed).shuffle(candidates)
    kept: List[str] = []
    seen_words = accepted = examined = 0
    for k, s in enumerate(candidates):
        ar = [m.group() for m in TOKEN_RE.finditer(s) if is_arabic_word(m.group())]
        oks = [checker.ok(w) for w in ar]
        examined += 1
        seen_words += len(ar)
        accepted += sum(oks)
        if (len(ar) - sum(oks)) / max(len(ar), 1) > max_reject_ratio:
            continue
        kept.append(s)
        if progress and k % 100 == 0:
            progress(min(len(kept) / max_sentences, 1.0))
        if len(kept) >= max_sentences:
            break
    stats = {
        "candidate_sentences": len(candidates),
        "examined": examined,
        "kept": len(kept),
        "dictionary_acceptance": accepted / max(seen_words, 1),
    }
    return kept, stats


def calibrate(
    sentences: List[str],
    scorer,
    checker: WordChecker,
    percentile: float = 0.5,
    only_accepted_words: bool = True,
    progress: Optional[Callable[[float], None]] = None,
) -> CalibrationResult:
    """threshold = the score below which `percentile` % of correct words fall.

    By default only words accepted by the dictionary are used, because those are the
    words that must NOT be flagged by the context step (the others are flagged by the
    dictionary step regardless).
    """
    scores: List[float] = []
    for i, s in enumerate(sentences):
        toks = tokenize(s)
        sc = scorer.score_words([t.norm for t in toks])
        for j, t in enumerate(toks):
            if j in sc and (not only_accepted_words or checker.ok(t.norm)):
                scores.append(sc[j])
        if progress:
            progress((i + 1) / max(len(sentences), 1))
    arr = np.array(scores, dtype=float)
    if arr.size == 0:
        raise ValueError("No words could be scored. Check that the text is Arabic.")
    return CalibrationResult(
        threshold=float(np.percentile(arr, percentile)),
        percentile=percentile,
        scores=arr,
        n_sentences=len(sentences),
        n_words=int(arr.size),
    )


# --------------------------------------------------------------------------------------
# Domain profiles (profiles/<key>.json)
# --------------------------------------------------------------------------------------
def _slug(key: str) -> str:
    key = re.sub(r"[^\w\-]+", "_", key.strip(), flags=re.UNICODE).strip("_")
    return key or "profile"


def load_profiles(directory: Path = PROFILES_DIR) -> Dict[str, dict]:
    profiles: Dict[str, dict] = {}
    directory = Path(directory)
    if directory.exists():
        for f in sorted(directory.glob("*.json")):
            try:
                data = json.loads(f.read_text(encoding="utf-8"))
                data.setdefault("label", f.stem)
                data.setdefault("threshold", FALLBACK_THRESHOLD)
                data.setdefault("calibrated", False)
                profiles[f.stem] = data
            except (json.JSONDecodeError, OSError):
                continue
    if not profiles:
        profiles["general"] = {
            "label": "General MSA (uncalibrated default)",
            "threshold": FALLBACK_THRESHOLD,
            "calibrated": False,
            "percentile": None,
            "n_words": 0,
            "note": "Starting guess, not measured. Calibrate on texts from your domain.",
        }
    return profiles


def save_profile(key: str, data: dict, directory: Path = PROFILES_DIR) -> str:
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    slug = _slug(key)
    (directory / f"{slug}.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return slug
